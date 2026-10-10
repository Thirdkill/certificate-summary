"""
Bonus Feature: Excel Export

Exports one or many extracted certificates into a nicely formatted
.xlsx spreadsheet using openpyxl, so results can be reviewed or
imported into other systems (e.g. a student records database).
"""

import os
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EXPORT_DIR = "exports"
os.makedirs(EXPORT_DIR, exist_ok=True)

HEADERS = [
    "Certificate ID", "Candidate Name", "Certificate Title",
    "Organization", "Issue Date", "Certificate Number",
    "Grade/Score", "Duration", "OCR Confidence (%)", "Processed At",
]


def export_certificates_to_excel(certificates: List, filename: str = "certificates_export.xlsx") -> str:
    """
    Writes a list of Certificate ORM objects to a styled Excel workbook
    and returns the path to the generated file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Certificates"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, cert in enumerate(certificates, start=2):
        ws.cell(row=row_idx, column=1, value=cert.id[:8])
        ws.cell(row=row_idx, column=2, value=cert.candidate_name or "N/A")
        ws.cell(row=row_idx, column=3, value=cert.certificate_title or "N/A")
        ws.cell(row=row_idx, column=4, value=cert.organization_name or "N/A")
        ws.cell(row=row_idx, column=5, value=cert.issue_date or "N/A")
        ws.cell(row=row_idx, column=6, value=cert.certificate_number or "N/A")
        ws.cell(row=row_idx, column=7, value=cert.grade_score or "N/A")
        ws.cell(row=row_idx, column=8, value=cert.duration or "N/A")
        ws.cell(row=row_idx, column=9, value=cert.ocr_confidence or 0)
        ws.cell(row=row_idx, column=10, value=str(cert.created_at))

    for col_idx in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    out_path = os.path.join(EXPORT_DIR, filename)
    wb.save(out_path)
    return out_path
